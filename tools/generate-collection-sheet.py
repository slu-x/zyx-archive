# -*- coding: utf-8 -*-
# 根据 data/nav.json + data/events/*.json 生成一份"资料收集表" Excel，给管理员传到腾讯文档、
# 分享给粉丝填写用（每个大 tab 一个 sheet，已有内容标记"已导入"，方便以后只处理新增行）。
#
# 运行方式：cd tools && pip install openpyxl && python generate-collection-sheet.py
# 依赖：openpyxl（如果这台机器还没装，先 pip install openpyxl）
import json
import os
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
NAV_PATH = os.path.join(ROOT, "data", "nav.json")
EVENTS_DIR = os.path.join(ROOT, "data", "events")
OUT_PATH = os.path.join(ROOT, "Archive of Lay - 资料收集表.xlsx")

with open(NAV_PATH, encoding="utf-8") as f:
    nav = json.load(f)

PLATFORM_LIST = ["微博", "抖音", "小红书", "知乎", "贴吧", "Bilibili", "YouTube", "媒体", "其他"]
YEAR_LIST = [str(y) for y in range(2012, 2027)] + ["不确定/未知"]
YESNO_LIST = ["否", "是"]
STATUS_LIST = ["待处理", "已导入"]

HEADER = ["分类路径", "标题", "链接", "平台", "年份", "附注", "是否文件夹链接（是=需要展开整理里面所有链接）", "状态（管理员用，不用改）"]


def walk_paths(node, prefix, acc):
    if node.get("type") in ("event", "link", "divider"):
        return
    title = node.get("title", node.get("id", ""))
    path = prefix + [title] if title else prefix
    acc.append(" > ".join(path))
    for c in node.get("children", []) or []:
        walk_paths(c, path, acc)


def walk_rows(node, prefix, tab_title, rows, events_cache):
    ntype = node.get("type")
    if ntype == "link":
        rows.append({
            "path": " > ".join(prefix) if prefix else tab_title,
            "title": node.get("title", ""),
            "url": node.get("url", ""),
            "platform": node.get("platform", ""),
            "year": "",
            "note": "",
        })
        return
    if ntype == "event":
        eid = node["eventId"]
        if eid not in events_cache:
            with open(os.path.join(EVENTS_DIR, f"{eid}.json"), encoding="utf-8") as f:
                events_cache[eid] = json.load(f)
        ev = events_cache[eid]
        year = (ev.get("date") or "")[:4]
        for m in ev.get("materials", []):
            rows.append({
                "path": " > ".join(prefix) if prefix else tab_title,
                "title": "[" + ev["title"] + "] " + m.get("title", ""),
                "url": m.get("url", ""),
                "platform": m.get("platform", ""),
                "year": year,
                "note": "事件：" + ev["title"] + "｜类别：" + m.get("category", ""),
            })
        return
    if ntype == "divider":
        return
    title = node.get("title", "")
    new_prefix = prefix + [title] if title else prefix
    for c in node.get("children", []) or []:
        walk_rows(c, new_prefix, tab_title, rows, events_cache)


wb = openpyxl.Workbook()
wb.remove(wb.active)

opt_ws = wb.create_sheet("选项来源")
opt_ws.sheet_state = "hidden"

col = 1
opt_ws.cell(row=1, column=col, value="平台")
for i, v in enumerate(PLATFORM_LIST):
    opt_ws.cell(row=i + 2, column=col, value=v)
platform_col_letter = get_column_letter(col)
platform_range = "选项来源!$" + platform_col_letter + "$2:$" + platform_col_letter + "$" + str(len(PLATFORM_LIST) + 1)
col += 1

opt_ws.cell(row=1, column=col, value="年份")
for i, v in enumerate(YEAR_LIST):
    opt_ws.cell(row=i + 2, column=col, value=v)
year_col_letter = get_column_letter(col)
year_range = "选项来源!$" + year_col_letter + "$2:$" + year_col_letter + "$" + str(len(YEAR_LIST) + 1)
col += 1

opt_ws.cell(row=1, column=col, value="是否文件夹")
for i, v in enumerate(YESNO_LIST):
    opt_ws.cell(row=i + 2, column=col, value=v)
yesno_col_letter = get_column_letter(col)
yesno_range = "选项来源!$" + yesno_col_letter + "$2:$" + yesno_col_letter + "$" + str(len(YESNO_LIST) + 1)
col += 1

opt_ws.cell(row=1, column=col, value="状态")
for i, v in enumerate(STATUS_LIST):
    opt_ws.cell(row=i + 2, column=col, value=v)
status_col_letter = get_column_letter(col)
status_range = "选项来源!$" + status_col_letter + "$2:$" + status_col_letter + "$" + str(len(STATUS_LIST) + 1)
col += 1

tabs = [t for t in nav["tabs"] if t.get("id") != "tribute"]

path_ranges = {}
for tab in tabs:
    paths = []
    walk_paths(tab, [], paths)
    header_col = col
    opt_ws.cell(row=1, column=header_col, value="分类路径__" + tab["id"])
    for i, v in enumerate(paths):
        opt_ws.cell(row=i + 2, column=header_col, value=v)
    letter = get_column_letter(header_col)
    path_ranges[tab["id"]] = "选项来源!$" + letter + "$2:$" + letter + "$" + str(len(paths) + 1)
    col += 1

HEADER_FILL = PatternFill(start_color="3D3267", end_color="3D3267", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

for tab in tabs:
    ws = wb.create_sheet(tab["title"][:31])
    for c_i, h in enumerate(HEADER, start=1):
        cell = ws.cell(row=1, column=c_i, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 10
    ws.column_dimensions["E"].width = 12
    ws.column_dimensions["F"].width = 30
    ws.column_dimensions["G"].width = 20
    ws.column_dimensions["H"].width = 14
    ws.freeze_panes = "A2"

    rows = []
    events_cache = {}
    for c in tab.get("children", []) or []:
        walk_rows(c, [], tab["title"], rows, events_cache)

    r = 2
    for row in rows:
        ws.cell(row=r, column=1, value=row["path"])
        ws.cell(row=r, column=2, value=row["title"])
        ws.cell(row=r, column=3, value=row["url"])
        ws.cell(row=r, column=4, value=row["platform"])
        ws.cell(row=r, column=5, value=row["year"])
        ws.cell(row=r, column=6, value=row["note"])
        ws.cell(row=r, column=7, value="否")
        ws.cell(row=r, column=8, value="已导入")
        r += 1

    last_data_row = max(r - 1, 1)
    last_row_with_buffer = last_data_row + 300

    dv_path = DataValidation(type="list", formula1="=" + path_ranges[tab["id"]], allow_blank=True, showDropDown=False)
    dv_path.error = "请从下拉列表选择一个已有分类；如果确实需要新分类，请在附注里写清楚，导入时会跟你确认。"
    dv_path.errorTitle = "分类不存在"
    ws.add_data_validation(dv_path)
    dv_path.add("A2:A" + str(last_row_with_buffer))

    dv_platform = DataValidation(type="list", formula1="=" + platform_range, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_platform)
    dv_platform.add("D2:D" + str(last_row_with_buffer))

    dv_year = DataValidation(type="list", formula1="=" + year_range, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_year)
    dv_year.add("E2:E" + str(last_row_with_buffer))

    dv_yesno = DataValidation(type="list", formula1="=" + yesno_range, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_yesno)
    dv_yesno.add("G2:G" + str(last_row_with_buffer))

    dv_status = DataValidation(type="list", formula1="=" + status_range, allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv_status)
    dv_status.add("H2:H" + str(last_row_with_buffer))

toc_ws = wb.create_sheet("目录", 0)
toc_ws.column_dimensions["A"].width = 22
toc_ws.column_dimensions["B"].width = 46
toc_ws.column_dimensions["C"].width = 12
toc_ws.cell(row=1, column=1, value="Sheet").font = Font(bold=True)
toc_ws.cell(row=1, column=2, value="包含的子分类（预览）").font = Font(bold=True)
toc_ws.cell(row=1, column=3, value="现有条目数").font = Font(bold=True)
for c in toc_ws[1]:
    c.fill = HEADER_FILL
    c.font = Font(color="FFFFFF", bold=True)
toc_ws.freeze_panes = "A2"

toc_row = 2
link_cell = toc_ws.cell(row=toc_row, column=1, value="填写说明")
link_cell.hyperlink = "#'填写说明'!A1"
link_cell.style = "Hyperlink"
toc_ws.cell(row=toc_row, column=2, value="怎么填这张表、下拉菜单怎么用")
toc_row += 1

for tab in tabs:
    sheet_name = tab["title"][:31]
    ws = wb[sheet_name]
    link_cell = toc_ws.cell(row=toc_row, column=1, value=sheet_name)
    link_cell.hyperlink = "#'" + sheet_name + "'!A1"
    link_cell.style = "Hyperlink"
    direct_children = [c.get("title") for c in (tab.get("children") or []) if c.get("title")]
    preview = "，".join(direct_children[:6]) + ("…" if len(direct_children) > 6 else "")
    toc_ws.cell(row=toc_row, column=2, value=preview or "（暂无子分类）")
    toc_ws.cell(row=toc_row, column=3, value=ws.max_row - 1)
    toc_row += 1

help_ws = wb.create_sheet("填写说明", 1)
help_ws.column_dimensions["A"].width = 100
lines = [
    "使用说明（给粉丝 & 管理员看）",
    "",
    "1. 每个大分类一个 sheet。",
    "2. 想加新资料，就在对应 sheet 底部空行开始填：分类路径、标题、链接。平台/年份/附注可以留空，管理员导入时会尽量补全。",
    "3. “分类路径”是下拉菜单，只能选表里已有的小分类；如果确实是全新的分类，没有合适的选项，就在“附注”里写清楚你想叫什么名字、放在哪个大分类下面，导入的时候会单独确认。",
    "4. “是否文件夹链接”：如果你要加的不是单条视频/文章，而是一整个 up 主合集、专栏、文件夹页面，选“是”，导入时会把里面能找到的链接都整理进来，重复的会自动去掉，不用你自己一条条复制。",
    "5. “状态”这一列是管理员用来标记“这条是不是已经处理过了”，不用管，看到已经写了内容的不用动它。",
    "6. 不确定的地方不用纠结，能填多少填多少，剩下的交给管理员导入时确认就行。",
]
for i, l in enumerate(lines, start=1):
    help_ws.cell(row=i, column=1, value=l)
help_ws["A1"].font = Font(bold=True, size=14)

wb.save(OUT_PATH)
print("saved:", OUT_PATH)
